# -*- coding: utf-8 -*-
import base64
import csv
import io
import unicodedata

from odoo import fields, models, _
from odoo.exceptions import UserError


class SystoreSalesCostImportWizard(models.TransientModel):
    _name = 'systore.sales.cost.import.wizard'
    _description = 'Importar costos manuales - Analítica de ventas'

    file_data = fields.Binary(string='Archivo', required=True)
    filename = fields.Char(string='Nombre del archivo')

    def _normalize_header(self, value):
        value = (value or '').strip().lower()
        value = ''.join(
            char for char in unicodedata.normalize('NFD', value)
            if unicodedata.category(char) != 'Mn'
        )
        return value.replace('_', ' ').replace('-', ' ').strip()

    def _read_rows(self):
        self.ensure_one()
        if not self.file_data:
            raise UserError(_('Adjunta un archivo Excel (.xlsx) o CSV.'))
        raw = base64.b64decode(self.file_data)
        name = (self.filename or '').lower()

        if name.endswith('.xlsx'):
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise UserError(_('El servidor no tiene disponible openpyxl para leer archivos .xlsx. Usa CSV como alternativa.')) from exc
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
        elif name.endswith('.csv') or not name:
            text = raw.decode('utf-8-sig')
            sample = text[:4096]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            rows = list(csv.reader(io.StringIO(text), dialect))
        else:
            raise UserError(_('Formato no soportado. Usa un archivo .xlsx o .csv.'))

        if not rows:
            raise UserError(_('El archivo está vacío.'))
        return rows

    def action_import(self):
        self.ensure_one()
        if not self.env.user.has_group('systore_sales_cost_analytics.group_systore_analytics_manager'):
            raise UserError(_('Solo un Administrador de Analítica de ventas puede importar costos.'))

        rows = self._read_rows()
        headers = [self._normalize_header(str(value or '')) for value in rows[0]]

        id_aliases = {'id', 'id linea', 'id de linea', 'line id', 'database id'}
        cost_aliases = {'costo unitario', 'costo', 'unit cost', 'unit cost company', 'unit cost company id', 'unit cost company'}

        id_idx = next((i for i, h in enumerate(headers) if h in id_aliases), None)
        cost_idx = next((i for i, h in enumerate(headers) if h in cost_aliases), None)

        if id_idx is None or cost_idx is None:
            raise UserError(_(
                'El archivo debe contener las columnas “ID línea” (o “ID”) y “Costo unitario”. '
                'Encabezados detectados: %s'
            ) % ', '.join(headers))

        Line = self.env['systore.sales.cost.line']
        updated = 0
        skipped = 0
        errors = []

        for row_number, row in enumerate(rows[1:], start=2):
            if not row or all(value in (None, '') for value in row):
                continue
            try:
                line_id_raw = row[id_idx] if id_idx < len(row) else None
                cost_raw = row[cost_idx] if cost_idx < len(row) else None
                if line_id_raw in (None, '') or cost_raw in (None, ''):
                    skipped += 1
                    continue
                line_id = int(float(line_id_raw))
                if isinstance(cost_raw, str):
                    normalized_cost = cost_raw.replace('$', '').replace(' ', '').replace(',', '')
                    cost = float(normalized_cost)
                else:
                    cost = float(cost_raw)
                if cost <= 0:
                    raise ValueError(_('El costo debe ser mayor que cero.'))

                line = Line.browse(line_id).exists()
                if not line:
                    raise ValueError(_('No existe la línea con ID %s.') % line_id)
                if line.purchase_order_line_id:
                    skipped += 1
                    continue
                if line.unit_cost_company and line.cost_source == 'purchase':
                    skipped += 1
                    continue

                line.write({'unit_cost_company': cost})
                updated += 1
            except Exception as exc:
                errors.append(_('Fila %(row)s: %(error)s') % {'row': row_number, 'error': str(exc)})

        if errors:
            preview = '\n'.join(errors[:10])
            if len(errors) > 10:
                preview += _('\n… y %s errores más.') % (len(errors) - 10)
            raise UserError(_(
                'Se encontraron errores en el archivo. No se completó la importación.\n\n%s'
            ) % preview)

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Importación de costos completada'),
                'message': _('%(updated)s líneas actualizadas; %(skipped)s omitidas.') % {
                    'updated': updated,
                    'skipped': skipped,
                },
                'type': 'success',
                'sticky': False,
                'next': {'type': 'ir.actions.act_window_close'},
            },
        }
